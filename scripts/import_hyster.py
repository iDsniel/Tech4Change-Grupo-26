"""Consolida exportações atuais do Hyster Tracker no contrato operacional Pulso v4.

Uso:
  python scripts/import_hyster.py /caminho/exports \
    --workforce /caminho/workforceKPITier7.xlsx \
    --output .data/Pulso-base-operacao.json

O importador preserva o grão original de cada fonte. Médias diárias/mensais do
Workforce continuam médias reportadas pela Hyster e nunca são transformadas em
uma série card-day artificial. O código do cartão é preservado; nomes de
operadores não são persistidos no pacote operacional.
"""
import argparse
import hashlib
import json
import math
import re
from collections import Counter
from datetime import date, datetime, time
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel

FILES = [
    "assetListing",
    "DailyFleetUtilization",
    "utilizationKPITier7",
    "AssetOperatingHistory",
    "currentFleetStatus",
    "PMTrackerDW",
    "fuelUsageSummaryequipment",
    "costOfOperationTier7",
]

MONTHS = {name: i for i, name in enumerate(
    ["JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"], 1)}

# Workforce: colunas zero-based (daily, monthly, total) começam em 3 e avançam 3.
WORKFORCE_METRICS = [
    ("serviceHours", "Main Service Meter", "hours"),
    ("driveHours", "Drive Motor / Engine Runtime Meter", "hours"),
    ("hydraulicMeterHours", "Hydraulic Operation Meter", "hours"),
    ("tractionMeterHours", "Transmission/Traction Operation Meter", "hours"),
    ("distanceKm", "Odometer", "km"),
    ("monitoredHours", "Monitored Duration", "hours"),
    ("keyHours", "Key Switch 'On' Duration", "hours"),
    ("presenceHours", "Operator Presence Duration", "hours"),
    ("motionHours", "Moving Duration", "hours"),
    ("hydraulicHours", "Hydraulic Function Duration", "hours"),
    ("workHours", "Working Duration", "hours"),
    ("liftHours", "Lift Duration", "hours"),
    ("lowerHours", "Lower Duration", "hours"),
    ("auxiliaryHydraulicHours", "Auxillary Hydraulic Duration", "hours"),
    ("lowSpeedHours", "Low Speed Operation Duration", "hours"),
    ("mediumSpeedHours", "Medium Speed Operation Duration", "hours"),
    ("highSpeedHours", "High Speed Operation Duration", "hours"),
    ("lowLevelOverspeedHours", "Low Level Over-speed Duration", "hours"),
    ("highLevelOverspeedHours", "High Level Over-speed Duration", "hours"),
    ("reverseHours", "Reverse Gear Operation Duration", "hours"),
    ("forwardHours", "Forward Gear Operation Duration", "hours"),
    ("seatBeltViolationHours", "Seat Belt Violation Duration", "hours"),
    ("idleHours", "Idle Time", "hours"),
    ("containerCount", "Container Count", "count"),
    ("energyFuelUsedLiters", "Energy/Fuel Used", "litres"),
    ("ladenHours", "Laden Hours", "hours"),
    ("unladenHours", "Un-Laden Hours", "hours"),
    ("workingUnladenHours", "Working Hours Un-Laden", "hours"),
    ("unladenDurationHours", "Un-Laden Duration", "hours"),
]

# Fleet KPI: (key, source label, unit, monthly-column-index-zero-based)
FLEET_KPI_METRICS = [
    ("serviceHours", "Medidor de serviço principal", "hours", 7),
    ("monitoredHours", "Duração monitorada", "hours", 10),
    ("keyHours", "Interruptor de chave 'Na' duração", "hours", 13),
    ("presenceHours", "Duração de presença do operador", "hours", 16),
    ("motionHours", "Duração em movimento", "hours", 19),
    ("hydraulicHours", "Duração da função hidráulica", "hours", 22),
    ("workHours", "Duração do trabalho", "hours", 25),
    ("liftHours", "Duração de elevador", "hours", 28),
    ("idleHours", "Tempo ocioso", "hours", 31),
    ("ladenHours", "Horas de carga", "hours", 34),
    ("unladenHours", "Un-carga horas", "hours", 37),
    ("workingUnladenHours", "Un-carga em movimento horas", "hours", 40),
    ("unladenDurationHours", "Duração un-carga", "hours", 43),
]


def non_negative(value):
    if value is None or value == "":
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) and number >= 0 else None


def percent_value(value):
    if value is None or value == "":
        return None
    if isinstance(value, str):
        text = value.strip().replace("%", "").replace(",", ".")
        return non_negative(text)
    number = non_negative(value)
    if number is None:
        return None
    return number * 100 if 0 <= number <= 1 else number


def card_code(value):
    if value is None or value == "":
        return None
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, (int, float)) and float(value).is_integer():
        return str(int(value))
    return str(value).strip()


def iso_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)):
        converted = from_excel(value)
        return converted.date().isoformat() if isinstance(converted, datetime) else converted.isoformat()
    text = str(value).strip()
    for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError(f"Data Hyster inválida: {value}")


def iso_time(value):
    if isinstance(value, datetime):
        return value.time().replace(microsecond=0).isoformat()
    if isinstance(value, time):
        return value.replace(microsecond=0).isoformat()
    text = str(value or "").strip()
    for fmt in ("%H:%M:%S", "%H:%M"):
        try:
            return datetime.strptime(text, fmt).time().isoformat()
        except ValueError:
            pass
    return text


def iso_datetime_text(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.replace(second=0, microsecond=0).isoformat(timespec="minutes")
    text = str(value).strip()
    for fmt in ("%b %d %Y %I:%M%p", "%b %d %Y  %I:%M%p", "%Y-%m-%dT%H:%M", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).isoformat(timespec="minutes")
        except ValueError:
            pass
    return text


def join_datetime(date_value, time_value):
    return f"{iso_date(date_value)}T{iso_time(time_value)}"


def duration_seconds(value):
    if value in (None, ""):
        return None
    parts = str(value).strip().split(":")
    if len(parts) != 3:
        return None
    try:
        h, m, s = [int(part) for part in parts]
        return h * 3600 + m * 60 + s
    except ValueError:
        return None


def source_metadata(path, book):
    return {
        "file": path.name,
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "sheets": [{
            "name": sheet.title,
            "range": sheet.calculate_dimension(),
            "rows": sheet.max_row,
            "state": sheet.sheet_state,
        } for sheet in book.worksheets],
    }


def parse_workforce_period(text):
    match = re.fullmatch(
        r"\s*(\d{2})\s+([A-Z]{3})\s+(\d{4})\s+\d{2}:\d{2}:\d{2}\s*-\s*"
        r"(\d{2})\s+([A-Z]{3})\s+(\d{4})\s+\d{2}:\d{2}:\d{2}\s*", str(text or ""))
    if not match or match.group(2) not in MONTHS or match.group(5) not in MONTHS:
        raise ValueError("Unsupported Workforce KPI period")
    start = f"{match.group(3)}-{MONTHS[match.group(2)]:02d}-{int(match.group(1)):02d}"
    end = f"{match.group(6)}-{MONTHS[match.group(5)]:02d}-{int(match.group(4)):02d}"
    return start, end


def metric_triplet(row, start, source_label, unit, workforce_order=True):
    first = non_negative(row[start] if start < len(row) else None)
    second = non_negative(row[start + 1] if start + 1 < len(row) else None)
    total = non_negative(row[start + 2] if start + 2 < len(row) else None)
    if workforce_order:
        daily, monthly = first, second
    else:
        monthly, daily = first, second
    if daily is None and monthly is None and total is None:
        return None
    return {
        "sourceLabel": source_label,
        "unit": unit,
        "dailyAverage": daily or 0,
        "monthlyAverage": monthly or 0,
        "total": total or 0,
    }


def workforce_values(row):
    metrics, statistics = {}, {}
    for index, (key, label, unit) in enumerate(WORKFORCE_METRICS):
        stat = metric_triplet(row, 3 + index * 3, label, unit, True)
        if stat is not None:
            statistics[key] = stat
            metrics[key] = stat["total"]
    return metrics, statistics


def convert_workforce(path, product_meta, expected_start, expected_end):
    book = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        required = {"Main Page", "Workforce KPI Report"}
        if not required.issubset(book.sheetnames):
            raise ValueError("Unsupported Workforce KPI workbook: required sheets not found")
        main = book["Main Page"]
        report = book["Workforce KPI Report"]
        start, end = parse_workforce_period(main.cell(13, 1).value)
        if (start, end) != (expected_start, expected_end):
            raise ValueError(f"Workforce KPI period {start}..{end} differs from Hyster period {expected_start}..{expected_end}")
        if str(main.cell(15, 1).value).strip() != "Operator" or str(main.cell(19, 1).value).strip() != "Metric":
            raise ValueError("Unsupported Workforce KPI configuration")

        for index, (_, expected_label, _) in enumerate(WORKFORCE_METRICS):
            start_col = 4 + index * 3
            if report.cell(1, start_col).value != expected_label:
                raise ValueError(f"Unsupported Workforce layout at column {start_col}: {report.cell(1, start_col).value!r}")
            if not str(report.cell(2, start_col).value or "").startswith("Daily Averages"):
                raise ValueError(f"Missing Workforce daily average at column {start_col}")
            if not str(report.cell(2, start_col + 1).value or "").startswith("Monthly Averages"):
                raise ValueError(f"Missing Workforce monthly average at column {start_col + 1}")
            if not str(report.cell(2, start_col + 2).value or "").startswith("Total Usage"):
                raise ValueError(f"Missing Workforce total at column {start_col + 2}")

        cards, current, by_code, warnings = [], None, {}, []
        for source_row, row in enumerate(report.iter_rows(min_row=3, values_only=True), 3):
            label = row[1] if len(row) > 1 else None
            if label in (None, ""):
                continue
            label = str(label)
            if "Product ID :" in label:
                if current is None:
                    warnings.append(f"Linha {source_row}: equipamento sem cartão pai; registro ignorado")
                    continue
                product_match = re.search(r"Product ID\s*:\s*(\d+)", label)
                product_id = int(product_match.group(1)) if product_match else None
                meta = product_meta.get(product_id)
                if not meta:
                    warnings.append(f"Linha {source_row}: equipamento não conciliado; registro ignorado")
                    continue
                if any(item["assetId"] == meta["assetId"] for item in current["assets"]):
                    raise ValueError(f"Duplicate Workforce asset {meta['assetId']} under source row {current['sourceRow']}")
                count = non_negative(row[2] if len(row) > 2 else None)
                if count is None or not float(count).is_integer():
                    raise ValueError(f"Invalid Workforce usage count at row {source_row}")
                metrics, statistics = workforce_values(row)
                current["assets"].append({
                    "assetId": meta["assetId"],
                    "productId": meta["productId"],
                    "serialNumber": meta["serialNumber"],
                    "trackerAssetId": meta["trackerAssetId"],
                    "serviceId": meta["serviceId"],
                    "equipmentName": meta["equipmentName"],
                    "usageCount": int(count),
                    "metrics": metrics,
                    "statistics": statistics,
                    "sourceRow": source_row,
                })
                continue

            match = re.search(r"\(([^()]*)\)\s*$", label)
            raw_code = match.group(1).strip() if match else None
            code = raw_code if raw_code and re.fullmatch(r"\d{1,64}", raw_code) else None
            quality = "complete" if code else "incomplete"
            if code in by_code:
                quality = "ambiguous"
                for previous in cards:
                    if previous["cardCode"] == code:
                        previous["cardQuality"] = "ambiguous"
            count = non_negative(row[2] if len(row) > 2 else None)
            if count is None or not float(count).is_integer():
                raise ValueError(f"Invalid Workforce usage count at row {source_row}")
            metrics, statistics = workforce_values(row)
            current = {
                "cardCode": code,
                "cardQuality": quality,
                "usageCount": int(count),
                "metrics": metrics,
                "statistics": statistics,
                "assets": [],
                "sourceRow": source_row,
            }
            cards.append(current)
            if code:
                by_code.setdefault(code, []).append(source_row)

        if not cards:
            raise ValueError("Workforce KPI contains no card rows")
        if any(card["cardQuality"] == "incomplete" for card in cards):
            warnings.append("Há linhas de cartão sem código completo; elas permanecem sem identificação no contrato.")
        if any(card["cardQuality"] == "ambiguous" for card in cards):
            warnings.append("Há códigos de cartão repetidos; as linhas foram sinalizadas como ambíguas e não devem ser agregadas automaticamente.")
        warnings.extend([
            "Daily and monthly averages are source-reported and must not be interpreted as a real card-by-day time series.",
            "Os totais por cartão permanecem na granularidade card-period e não são distribuídos artificialmente por mês ou dia.",
            "Linhas filhas por equipamento ficam aninhadas no cartão e não são somadas novamente ao total do cartão.",
        ])

        availability = {}
        for key, label, unit in WORKFORCE_METRICS:
            values = [card["metrics"].get(key, 0) for card in cards]
            availability[key] = {
                "sourceLabel": label,
                "unit": unit,
                "cardsWithNonZero": sum(1 for value in values if value > 0),
                "cardsTotal": len(cards),
                "totalAcrossCards": round(sum(values), 6),
            }

        return ({
            "periodStart": start,
            "periodEnd": end,
            "granularity": "card-period",
            "unitSystem": "metric",
            "sourceFile": path.name,
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "cards": cards,
            "metricAvailability": availability,
            "warnings": warnings,
        }, source_metadata(path, book))
    finally:
        book.close()


def load_books(folder):
    books, rows, sources = {}, {}, []
    for name in FILES:
        path = folder / f"{name}.xlsx"
        if not path.exists():
            raise FileNotFoundError(path)
        book = openpyxl.load_workbook(path, data_only=True)
        books[name] = book
        rows[name] = [list(sheet.values) for sheet in book.worksheets]
        sources.append(source_metadata(path, book))
    return books, rows, sources


def parse_assets(rows):
    assets, product_meta = [], {}
    for source_row, r in enumerate(rows["assetListing"][1][1:], 2):
        if not r or r[2] in (None, ""):
            continue
        product_id = int(r[2])
        tracker_id = str(r[5])
        asset_id = "EP" + tracker_id.zfill(2)
        meta = {
            "assetId": asset_id,
            "trackerAssetId": tracker_id,
            "serviceId": str(r[6]),
            "productId": product_id,
            "equipmentName": str(r[3]),
            "serialNumber": str(r[4]),
            "site": str(r[0]),
            "department": str(r[1]),
            "serviceMeterHours": non_negative(r[7]) or 0,
            "sourceRow": source_row,
        }
        assets.append(meta)
        product_meta[product_id] = meta
    return assets, product_meta


def convert(folder, workforce=None):
    books, rows, sources = load_books(folder)
    try:
        assets, product_meta = parse_assets(rows)
        product_to_asset = {product_id: meta["assetId"] for product_id, meta in product_meta.items()}

        daily, seen, current_asset = [], set(), None
        for source_row, r in enumerate(rows["DailyFleetUtilization"][1][3:], 4):
            if len(r) < 12 or r[1] in (None, ""):
                continue
            if r[0]:
                match = re.search(r"ID de produto\s*:\s*(\d+)", str(r[0]))
                if not match or int(match.group(1)) not in product_to_asset:
                    raise ValueError(f"Daily Fleet asset not reconciled at row {source_row}")
                current_asset = product_to_asset[int(match.group(1))]
            if current_asset is None:
                raise ValueError(f"Daily Fleet row without asset at row {source_row}")
            day = iso_date(r[1])
            if (current_asset, day) in seen:
                raise ValueError(f"Duplicate daily record at row {source_row}")
            seen.add((current_asset, day))
            daily.append({
                "assetId": current_asset,
                "date": day,
                "keyHours": non_negative(r[2]) or 0,
                "presenceHours": non_negative(r[3]) or 0,
                "workHours": non_negative(r[4]) or 0,
                "idleHours": non_negative(r[5]) or 0,
                "equipmentUsedCount": int(non_negative(r[6]) or 0),
                "usedPercent": percent_value(r[7]) or 0,
                "workPercent": percent_value(r[8]) or 0,
                "idlePercent": percent_value(r[9]) or 0,
                "waitHours": non_negative(r[10]) or 0,
                "waitPercentOfIdle": percent_value(r[11]) or 0,
                "sourceRow": source_row,
            })

        dates = [row["date"] for row in daily]
        if not dates:
            raise ValueError("Daily Fleet contains no data")
        period_start, period_end = min(dates), max(dates)

        kpi = []
        for source_row, r in enumerate(rows["utilizationKPITier7"][1][2:], 3):
            if len(r) < 46 or r[2] in (None, ""):
                continue
            product_id = int(r[2])
            if product_id not in product_meta:
                raise ValueError(f"Fleet KPI asset not reconciled at row {source_row}")
            metrics, statistics = {}, {}
            for key, label, unit, start in FLEET_KPI_METRICS:
                stat = metric_triplet(r, start, label, unit, False)
                if stat is not None:
                    statistics[key] = stat
                    metrics[key] = stat["total"]
            item = {
                "assetId": product_meta[product_id]["assetId"],
                **metrics,
                "loadHours": metrics.get("ladenHours", 0),
                "statistics": statistics,
                "sourceRow": source_row,
            }
            kpi.append(item)

        events = []
        for source_row, r in enumerate(rows["AssetOperatingHistory"][1][1:], 2):
            if len(r) < 18 or r[1] in (None, "") or r[8] in (None, ""):
                continue
            product_id = int(r[1])
            if product_id not in product_meta:
                continue
            event_date = iso_date(r[8])
            if event_date < period_start or event_date > period_end:
                continue
            event_time = iso_time(r[9])
            start_date = iso_date(r[6])
            start_time = iso_time(r[7])
            events.append({
                "assetId": product_meta[product_id]["assetId"],
                "productId": product_id,
                "equipmentName": str(r[2]),
                "serialNumber": str(r[3]),
                "trackerAssetId": str(r[4]),
                "serviceId": str(r[5]),
                "date": event_date,
                "time": event_time,
                "eventAt": f"{event_date}T{event_time}",
                "startDate": start_date,
                "startTime": start_time,
                "startAt": f"{start_date}T{start_time}",
                "cardCode": card_code(r[12]),
                "sourceCritical": str(r[13]).strip().lower() in {"sim", "yes", "true"},
                "lockout": str(r[14]).strip().lower() in {"sim", "yes", "true"},
                "shutdown": str(r[15]).strip().lower() in {"sim", "yes", "true"},
                "sourceStatus": str(r[16]),
                "type": str(r[17]),
                "sourceRow": source_row,
            })

        status_book = books["currentFleetStatus"]
        status_rows = rows["currentFleetStatus"]
        snapshot_at = iso_datetime_text(status_book["Folha1"].cell(19, 2).value)
        current_status = []
        for source_row, r in enumerate(status_rows[2][2:], 3):
            if len(r) < 11 or r[1] in (None, ""):
                continue
            product_id = int(r[1])
            if product_id not in product_meta:
                continue
            equipment_class = str(r[5]) if r[5] not in (None, "") else None
            product_meta[product_id]["equipmentClass"] = equipment_class
            current_status.append({
                "assetId": product_meta[product_id]["assetId"],
                "productId": product_id,
                "equipmentName": str(r[2]),
                "serialNumber": str(r[3]),
                "trackerAssetId": str(r[4]),
                "equipmentClass": equipment_class,
                "status": str(r[6]),
                "lockoutCause": None if r[7] in (None, "") else str(r[7]),
                "lastAccessedAt": iso_datetime_text(r[9]),
                "driveDuration": None if r[10] in (None, "") else str(r[10]),
                "driveDurationSeconds": duration_seconds(r[10]),
                "sourceRow": source_row,
            })

        # Reflect class back into the public asset catalog without operator identity.
        for asset in assets:
            meta = product_meta[asset["productId"]]
            if meta.get("equipmentClass"):
                asset["equipmentClass"] = meta["equipmentClass"]

        fuel = []
        for source_row, r in enumerate(rows["fuelUsageSummaryequipment"][1][1:], 2):
            if len(r) < 11 or r[3] in (None, ""):
                continue
            product_id = int(r[3])
            if product_id not in product_meta:
                continue
            fuel.append({
                "assetId": product_meta[product_id]["assetId"],
                "productId": product_id,
                "equipmentName": str(r[4]),
                "serialNumber": str(r[5]),
                "trackerAssetId": str(r[6]),
                "serviceId": str(r[7]),
                "monthlyAverageLiters": non_negative(r[8]) or 0,
                "dailyAverageLiters": non_negative(r[9]) or 0,
                "reportedLiters": non_negative(r[10]) or 0,
                "sourceRow": source_row,
            })

        costs = []
        for source_row, r in enumerate(rows["costOfOperationTier7"][1][2:], 3):
            if len(r) < 12 or r[1] in (None, ""):
                continue
            product_id = int(r[1])
            if product_id not in product_meta:
                continue
            costs.append({
                "assetId": product_meta[product_id]["assetId"],
                "productId": product_id,
                "equipmentName": str(r[2]),
                "serialNumber": str(r[3]),
                "trackerAssetId": str(r[4]),
                "serviceId": str(r[5]),
                "startHours": non_negative(r[6]) or 0,
                "endHours": non_negative(r[7]) or 0,
                "intervalHours": non_negative(r[8]) or 0,
                "costDetails": r[9],
                "reportedCostPerHour": non_negative(r[10]) or 0,
                "reportedTotal": non_negative(r[11]),
                "sourceRow": source_row,
            })

        pm_sheet = books["PMTrackerDW"]["PM Tracker relatório"]
        pm_message = str(pm_sheet.cell(1, 1).value or "")
        maintenance_available = not pm_message.startswith("Nenhum PM Tracker Data")
        maintenance = {
            "available": maintenance_available,
            "message": pm_message or None,
            "periodStart": period_start,
            "periodEnd": period_end,
            "records": [],
        }

        event_type_counts = Counter(event["type"] for event in events)
        event_status_counts = Counter(event["sourceStatus"] for event in events)
        event_asset_counts = Counter(event["assetId"] for event in events)
        event_summary = {
            "totalEvents": len(events),
            "byType": dict(event_type_counts),
            "byStatus": dict(event_status_counts),
            "byAsset": dict(event_asset_counts),
            "cardsObserved": len({event["cardCode"] for event in events if event.get("cardCode")}),
        }
        fleet_status_summary = {
            "total": len(current_status),
            "byStatus": dict(Counter(item["status"] for item in current_status)),
        }

        site = assets[0].get("site") if assets else ""
        department = assets[0].get("department") if assets else ""
        equipment_classes = sorted({asset.get("equipmentClass") for asset in assets if asset.get("equipmentClass")})
        operation_id = hashlib.sha256(str(site).encode()).hexdigest()[:16]
        calendar_days = (date.fromisoformat(period_end) - date.fromisoformat(period_start)).days + 1
        possible_asset_days = calendar_days * len(assets)
        by_asset_daily = Counter(item["assetId"] for item in daily)

        result = {
            "schemaVersion": 4,
            "provider": "Hyster Tracker",
            "operationId": operation_id,
            "periodStart": period_start,
            "periodEnd": period_end,
            "granularity": "asset-day",
            "unitSystem": "metric",
            "operation": {
                "site": site,
                "department": department,
                "equipmentClass": equipment_classes[0] if len(equipment_classes) == 1 else None,
                "assetCount": len(assets),
            },
            "assets": sorted(assets, key=lambda item: item["assetId"]),
            "daily": daily,
            "kpi": kpi,
            "events": events,
            "eventSummary": event_summary,
            "currentStatusSnapshotAt": snapshot_at,
            "currentStatus": current_status,
            "fleetStatusSummary": fleet_status_summary,
            "fuel": fuel,
            "costs": costs,
            "maintenanceAvailable": maintenance_available,
            "maintenance": maintenance,
            "reportMetadata": {
                "dailyFleet": {"granularity": "asset-day", "sourceFile": "DailyFleetUtilization.xlsx"},
                "fleetKpi": {"granularity": "asset-period", "sourceFile": "utilizationKPITier7.xlsx"},
                "operatingHistory": {"granularity": "event", "sourceFile": "AssetOperatingHistory.xlsx", "criticalOnly": all(event["sourceCritical"] for event in events) if events else False},
                "currentFleetStatus": {"granularity": "snapshot", "sourceFile": "currentFleetStatus.xlsx", "snapshotAt": snapshot_at},
                "fuel": {"granularity": "asset-period", "sourceFile": "fuelUsageSummaryequipment.xlsx"},
                "cost": {"granularity": "asset-period", "sourceFile": "costOfOperationTier7.xlsx"},
                "maintenance": {"granularity": "asset-period", "sourceFile": "PMTrackerDW.xlsx"},
            },
            "dataQuality": {
                "dailyFleetCoverage": {
                    "calendarDays": calendar_days,
                    "assets": len(assets),
                    "possibleAssetDays": possible_asset_days,
                    "rowsPresent": len(daily),
                    "rowsOmitted": max(0, possible_asset_days - len(daily)),
                    "byAsset": dict(by_asset_daily),
                },
                "fuelAllZero": bool(fuel) and all(item["reportedLiters"] == 0 for item in fuel),
                "costRateAllZero": bool(costs) and all(item["reportedCostPerHour"] == 0 for item in costs),
                "costTotalsMissing": bool(costs) and all(item["reportedTotal"] is None for item in costs),
                "maintenanceAvailable": maintenance_available,
                "eventExportCriticalOnly": bool(events) and all(event["sourceCritical"] for event in events),
                "eventStartAndEventTimestampEqualForAllRows": bool(events) and all(event["startAt"] == event["eventAt"] for event in events),
                "currentStatusSnapshotOutsideAnalysisPeriod": bool(snapshot_at) and snapshot_at[:10] > period_end,
                "notes": [
                    "Linhas equipamento-dia ausentes permanecem ausentes; o Pulso não converte ausência em zero.",
                    "Percentuais reportados pelo Daily Fleet são preservados como fornecidos, mesmo quando ultrapassam 100%.",
                    "Status atual é snapshot e não é tratado como evidência histórica do período quando capturado fora da janela.",
                ],
            },
            "sources": sources,
        }

        if workforce:
            result["workforce"], workforce_source = convert_workforce(workforce, product_meta, period_start, period_end)
            result["sources"].append(workforce_source)
            availability = result["workforce"]["metricAvailability"]
            result["dataQuality"]["workforceMetricAvailability"] = availability
            result["dataQuality"]["workforceMetricsAllZero"] = [key for key, value in availability.items() if value["cardsWithNonZero"] == 0]
            result["reportMetadata"]["workforce"] = {
                "granularity": "card-period",
                "sourceFile": Path(workforce).name,
                "groupBy": "Operator card",
                "statistics": ["dailyAverage", "monthlyAverage", "total"],
                "averageSemantics": "Source-reported averages; denominator is not reconstructed by Pulso.",
            }
        return result
    finally:
        for book in books.values():
            book.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("folder", type=Path)
    parser.add_argument("--output", type=Path, default=Path(".data/Pulso-base-operacao.json"))
    parser.add_argument("--workforce", type=Path, help="Workforce KPI export; mantém médias nativas e totais no grão card-period")
    args = parser.parse_args()
    data = convert(args.folder, args.workforce)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(json.dumps({
        "output": str(args.output),
        "schemaVersion": data["schemaVersion"],
        "periodStart": data["periodStart"],
        "periodEnd": data["periodEnd"],
        "assetDays": len(data["daily"]),
        "events": len(data["events"]),
        "cards": len(data.get("workforce", {}).get("cards", [])),
    }, ensure_ascii=False))
