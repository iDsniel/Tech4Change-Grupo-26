"""Preserve historical KPI source rows without names or guessed card numbers."""
import hashlib
import re
import unicodedata
from collections import Counter

import openpyxl


def convert_legacy(path):
    book = openpyxl.load_workbook(path, data_only=True)
    raw = book["Relatório KPI de força de traba"]
    base = book["Base dados"]
    aux = book["Auxiliar Oficial"]
    mapping = {"uses": 8, "hydraulicMeterHours": 11, "tractionMeterHours": 14,
               "distanceKm": 17, "monitoredHours": 20, "keyHours": 23,
               "presenceHours": 26, "motionHours": 29, "hydraulicHours": 32,
               "workHours": 35, "liftHours": 38, "lowerHours": 41,
               "highSpeedHours": 44, "reverseHours": 47, "forwardHours": 50, "idleHours": 53}
    period = str(raw.cell(3, 1).value)
    def normalized(value):
        return " ".join(unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode().upper().split())
    official = {normalized(aux.cell(n, 2).value): n for n in range(6, 38)}
    rows = []
    for n in range(3, raw.max_row + 1):
        label = str(raw.cell(n, 2).value or "")
        if not label:
            continue
        match = re.search(r"\((\d+)\)\s*$", label)
        code = match.group(1) if match else None
        measures = {key: raw.cell(n, col).value for key, col in mapping.items()}
        measures = {key: v if isinstance(v, (int, float)) else None for key, v in measures.items()}
        # Preserve the actual formula/cache layer separately from raw counters.
        base_row = n - 1
        cached = {str(base.cell(1, c).value): base.cell(base_row, c).value for c in range(3, 37)} if base_row <= 33 else {}
        cached = {k: v if isinstance(v, (int, float)) else None for k, v in cached.items()}
        official_row = official.get(normalized(label.split("(")[0]))
        display = {str(aux.cell(5, c).value): aux.cell(official_row, c).value for c in range(3, 24)} if official_row else {}
        display = {k: v for k, v in display.items() if isinstance(v, (int, float)) or v in ("OGMO", "ELDORADO")}
        rows.append({"rowId": f"LEG-{n}", "cardCode": code, "cardQuality": "complete" if code else "incomplete",
                     "sourceRow": n, "metrics": measures, "legacyCalculated": cached, "legacyDisplay": display})
    counts = Counter(r["cardCode"] for r in rows if r["cardCode"])
    for row in rows:
        if counts[row["cardCode"]] > 1:
            row["cardQuality"] = "ambiguous"
    # Auxiliar Oficial contains manually maintained display data, not an event ledger.
    display_impacts = [aux.cell(n, 22).value for n in range(6, 38)]
    impacts = []
    for n in range(2, book["Impactos"].max_row + 1):
        r = [book["Impactos"].cell(n, c).value for c in range(1, 23)]
        impacts.append({"dateText": str(r[5]), "cardCode": str(r[10]), "type": str(r[11]), "sourceRow": n})
    result = {"periodLabel": period, "sourceFile": path.name,
              "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
              "rows": rows, "impactArchive": impacts,
              "displayImpactTotal": sum(v for v in display_impacts if isinstance(v, (int, float))),
              "displayFuelKg": book["Auxiliar Visual"].cell(5, 1).value,
              "sheets": [{"name": s.title, "state": s.sheet_state} for s in book],
              "warnings": ["Main Page indica setembro/2022; período KPI indica agosto/2023.",
                           "Impactos contém eventos de 2022; não somar aos impactos de agosto/2023.",
                           "Auxiliar Oficial contém valores estáticos e não reconcilia automaticamente com a base.",
                           "Cartões sem fechamento de parênteses estão incompletos; códigos repetidos são ambíguos.",
                           "Combustível em kg é um valor manual sem conciliação; não converter para litros.",
                           "Metadados de localidade divergem do nome do arquivo; histórico mantido separado da operação atual."]}
    book.close()
    return result
